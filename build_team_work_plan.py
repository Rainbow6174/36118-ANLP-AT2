from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "team_planning"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT_DIR / "AT2_Role_Selection_and_HD_Delivery_Notice.docx"


# compact_reference_guide preset tokens
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11
MARGIN_IN = 1.0
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5D6875"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GRID = "B8C4D1"
GREEN = "2E6B57"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=5, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"]
    if inside:
        names.extend(["insideH", "insideV"])
    for name in names:
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell, **CELL_MARGIN_DXA)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    set_paragraph_spacing(p, after=4, line=1.25)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.25)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=7, line=1.20)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.08)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=10.5, color=DARK_BLUE, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=NAVY)


def add_code_block(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=8, line=1.05)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=9.2, color=NAVY)
    return p


def style_cell_text(cell, header=False, size=9.4):
    for p in cell.paragraphs:
        set_paragraph_spacing(p, before=0, after=0, line=1.12)
        for run in p.runs:
            set_run_font(run, size=size, color=WHITE if header else NAVY, bold=header)


def add_summary_table(doc):
    headers = ["Member", "Workstream", "Best suited to", "Primary rubric impact"]
    rows = [
        ["1", "Data, coordination and assembly", "Data engineering, EDA, documentation and organised delivery", "A2, B1 and B3"],
        ["2", "Sentiment modelling", "Traditional ML, Transformer training and controlled experiments", "A2 and B1"],
        ["3", "Sarcasm modelling", "Imbalanced classification, thresholds and error analysis", "A2 and B1"],
        ["4", "Variety-aware innovation", "Experimental research, ablations and cross-domain reasoning", "A1, A2 and B1"],
        ["5", "Evaluation and analysis", "Statistics, reproducible metrics, charts and interpretation", "A2, A3, B1 and B3"],
        ["6", "Web application and UX", "Frontend development, integration, deployment and user testing", "A3 and B2"],
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [720, 2200, 3860, 2580])
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        header.cells[idx].text = value
        set_cell_shading(header.cells[idx], NAVY)
        style_cell_text(header.cells[idx], header=True, size=9.2)
    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            row.cells[idx].text = value
            if int(values[0]) % 2 == 0:
                set_cell_shading(row.cells[idx], LIGHT_GRAY)
            style_cell_text(row.cells[idx], size=9.1)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, [720, 2200, 3860, 2580])
    return table


def add_schema_table(doc):
    headers = ["Field", "Definition"]
    rows = [
        ["id", "Stable identifier shared by training, evaluation and application modules."],
        ["text", "Raw text sample used as model input."],
        ["variety", "English variety label: en-AU, en-IN or en-UK."],
        ["source", "Original source domain, such as Google or Reddit."],
        ["sentiment", "Gold sentiment label."],
        ["sarcasm", "Gold sarcasm label."],
        ["split", "Frozen train, validation or test assignment."],
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1800, 7560])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        style_cell_text(table.rows[0].cells[idx], header=True)
    for i, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            row.cells[idx].text = value
            if i % 2:
                set_cell_shading(row.cells[idx], LIGHT_GRAY)
            style_cell_text(row.cells[idx], size=9.5)
        set_run_font(row.cells[0].paragraphs[0].runs[0], name="Consolas", size=9.2, color=DARK_BLUE, bold=True)
    set_table_geometry(table, [1800, 7560])
    return table


def add_prediction_table(doc):
    headers = ["Field", "Purpose"]
    rows = [
        ["id", "Joins predictions back to the frozen test data."],
        ["true_label", "Gold label used for evaluation."],
        ["predicted_label", "Model decision."],
        ["probability", "Confidence score for calibration and threshold analysis."],
        ["model_name", "Human-readable model identifier."],
        ["model_version", "Version used to reproduce a reported result."],
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2200, 7160])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        style_cell_text(table.rows[0].cells[idx], header=True)
    for i, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            row.cells[idx].text = value
            if i % 2:
                set_cell_shading(row.cells[idx], LIGHT_GRAY)
            style_cell_text(row.cells[idx], size=9.5)
        set_run_font(row.cells[0].paragraphs[0].runs[0], name="Consolas", size=9.2, color=DARK_BLUE, bold=True)
    set_table_geometry(table, [2200, 7160])
    return table


def add_work_package(
    doc,
    number,
    title,
    purpose,
    best_fit,
    tasks,
    deliverables,
    definition_done,
    report_section,
    primary_scores,
    secondary_scores,
    failure_risk,
    bullet_id,
):
    add_heading(doc, f"Member {number}: {title}", 2)
    add_body(doc, f"Purpose: {purpose}", bold_prefix="Purpose:", after=4)
    add_body(doc, f"Best suited to: {best_fit}", bold_prefix="Best suited to:", after=4)
    p = add_heading(doc, "Core responsibilities", 3)
    for task in tasks:
        add_list_item(doc, task, bullet_id)
    p = add_heading(doc, "Required handoff", 3)
    add_body(doc, deliverables, after=4)
    p = add_heading(doc, "Definition of Done", 3)
    for item in definition_done:
        add_list_item(doc, item, bullet_id)
    add_body(doc, f"Report ownership: {report_section}", bold_prefix="Report ownership:", after=8)
    add_callout(
        doc,
        "SCORING IMPACT",
        f"Primary: {primary_scores}. Secondary: {secondary_scores}. If incomplete: {failure_risk}",
    )


def add_footer_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, before=0, after=0, line=1.0)
    label = p.add_run("AT2 Role Selection and Delivery Notice  |  ")
    set_run_font(label, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    r_pr.append(color)
    r_pr.append(size)
    r.append(r_pr)
    r.extend([begin, instr, separate, value, end])
    p._p.append(r)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(12.5)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer_page_number(section)

    props = doc.core_properties
    props.title = "AT2 Role Selection and HD-Oriented Delivery Notice"
    props.subject = "Low-dependency role selection, delivery standards and rubric impact for a six-person ANLP project team"
    props.author = "AT2 Project Team"
    props.keywords = "AT2, ANLP, English varieties, sentiment, sarcasm, team plan"

    bullet_id = add_numbering_definition(doc, "bullet")
    number_id = add_numbering_definition(doc, "decimal")

    kicker = doc.add_paragraph()
    set_paragraph_spacing(kicker, before=0, after=5, line=1.0)
    r = kicker.add_run("AT2 ROLE SELECTION NOTICE")
    set_run_font(r, size=10, color=GREEN, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("Role Selection and\nHD-Oriented Delivery Plan")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("VarietyLens - English variety-aware sentiment and sarcasm classification")

    meta = doc.add_paragraph()
    set_paragraph_spacing(meta, before=0, after=14, line=1.15)
    for label, value in (
        ("Project", "VarietyLens"),
        ("Team size", "Six members"),
        ("Working model", "Independent work packages with fixed handoff contracts"),
        ("Document status", "Working draft for team agreement"),
    ):
        lr = meta.add_run(f"{label}: ")
        set_run_font(lr, size=10.2, color=NAVY, bold=True)
        vr = meta.add_run(value)
        set_run_font(vr, size=10.2, color=MUTED)
        meta.add_run().add_break()

    add_callout(
        doc,
        "KEY PRINCIPLE",
        "Each member owns a separate module, directory and report section. Collaboration occurs only at agreed data, prediction and application interfaces. SPARKPlus records actual contribution; it does not create a compulsory peer-review task.",
    )

    add_heading(doc, "1. Project objective", 1)
    add_body(
        doc,
        "VarietyLens will be a web-based NLP application for analysing sentiment and sarcasm across Australian, Indian and British English. The project will compare traditional NLP baselines with Transformer-based approaches, examine cross-variety generalisation, and communicate model confidence and performance through an accessible interface.",
    )
    add_body(
        doc,
        "The allocation below is designed to minimise day-to-day interference. It avoids a fully sequential workflow in which one delayed task blocks the whole team. Instead, the team agrees on shared formats early and then works in parallel.",
    )

    add_heading(doc, "2. Full-mark-oriented project design", 1)
    add_callout(
        doc,
        "PLANNING POSITION",
        "The plan is designed to create evidence for every rubric category and every basic submission requirement. It cannot guarantee a mark: final scoring still depends on execution quality, real results and assessor judgement.",
    )

    add_heading(doc, "2.1 Primary user and user contribution", 2)
    add_body(
        doc,
        "Primary users are customer and social insights analysts who examine English-language reviews and public social-media feedback across Australia, India and the United Kingdom. VarietyLens converts individual or uploaded texts into sentiment, sarcasm and confidence outputs; compares generic and variety-aware predictions; highlights uncertain cases for human review; and communicates known performance differences across English varieties.",
    )
    add_body(
        doc,
        "The system is decision support, not an identity detector. Known market or user-selected English variety should be the core input. Automatic variety detection is optional and must not be presented as nationality inference.",
    )

    add_heading(doc, "2.2 Core research questions", 2)
    for question in [
        "RQ1. How do traditional NLP models and Transformer-based models compare on sentiment and sarcasm classification across English varieties?",
        "RQ2. How does model performance differ across Australian, Indian and British English and across source domains?",
        "RQ3. Does a variety-aware approach improve classification reliability compared with a variety-agnostic model?",
    ]:
        add_list_item(doc, question, number_id)

    add_heading(doc, "2.3 Rubric legend", 2)
    rubric_items = [
        "A1 - Poster: Project design and innovation (35 points).",
        "A2 - Poster: Technical execution and evaluation (35 points).",
        "A3 - Poster: Quality of communication and visual aids (30 points).",
        "B1 - Report/code: NLP quality, justification and reproducibility (40 points).",
        "B2 - Report/code: Design features, ease of use and communication of insights (40 points).",
        "B3 - Report/code: Report readability, structure and originality (20 points).",
    ]
    for item in rubric_items:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "2.4 Project-wide evidence required", 2)
    evidence_items = [
        "One precise user problem, one primary user group and a consistent value proposition across the application, report and poster.",
        "A frozen and documented BESSTIE dataset, official split where available, duplicate and leakage checks, label distributions and a data limitations statement.",
        "A model ladder consisting of a majority baseline, TF-IDF traditional models, a Transformer and one controlled variety-aware innovation.",
        "Macro-F1 as the primary metric, plus per-class, per-variety and per-source results, cross-variety evaluation, uncertainty or calibration, repeated runs or confidence intervals, ablation and structured error analysis.",
        "A web application supporting single-text and CSV batch analysis, market/variety control, sentiment, sarcasm, confidence, model comparison, human-review flags, filtering, dashboard insights and export.",
        "A real usability study, one documented interface-improvement cycle, a clean-environment reproducibility test and consistent final numbers across report, application and poster.",
    ]
    for item in evidence_items:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "2.5 Basic submission coverage", 2)
    for item in [
        "A functioning web-based NLP application for a specific problem and user group.",
        "Source code, data and model instructions sufficient for reproduction.",
        "A detailed technical report covering objectives, phases, contributions, data, methods and justification, findings, outcomes, challenges and solutions, recommendations, next steps and appendix material.",
        "A digital group poster in editable PowerPoint, Canva or equivalent source format.",
        "An individual submission by every member where required, an accurate contribution statement and completed SPARKPlus self/peer assessment.",
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "3. Team allocation at a glance", 1)
    add_summary_table(doc)
    note = doc.add_paragraph()
    set_paragraph_spacing(note, before=5, after=8, line=1.15)
    r = note.add_run("Note: ")
    set_run_font(r, size=9.5, color=MUTED, bold=True)
    r = note.add_run("Names should replace Member 1-6 after the team agrees on ownership. The table describes responsibility, not a fixed percentage of the final mark.")
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    add_heading(doc, "4. Independent work packages", 1)

    add_work_package(
        doc,
        1,
        "Data, coordination and final assembly",
        "Create a stable project foundation that other members can use without waiting for repeated data changes.",
        "Someone organised and careful with data engineering, exploratory analysis, documentation, repository structure and final report assembly. Workload is highest at the beginning and end of the project.",
        [
            "Acquire and document the BESSTIE dataset and its permitted use.",
            "Clean text, inspect duplicates and define one frozen train/validation/test split.",
            "Produce exploratory analysis for variety, source, sentiment and sarcasm labels.",
            "Publish dataset_v1 within the first two or three working days and change it only for a serious defect.",
            "Maintain the shared repository conventions and combine the final report sections without rewriting their technical meaning.",
        ],
        "Processed data files, a reusable data loader, EDA figures, a data card, the agreed repository structure and the assembled report.",
        [
            "dataset_v1 is frozen, versioned and used by all modelling streams.",
            "Duplicate, leakage, missing-value and label-distribution checks are documented.",
            "Cleaning preserves spelling, punctuation, emoji, colloquial language and other variety-relevant signals unless a controlled experiment justifies removal.",
            "README commands work from a clean environment without personal absolute paths or secrets.",
            "The final report contains all required sections, consistent terminology and an accurate contribution table.",
        ],
        "Project objectives and scope; data sources and preprocessing; project phases and contribution summary; final document assembly.",
        "A2 through valid data evidence, B1 through reproducibility, and B3 through coherent report assembly",
        "A1 and A3 through clear project framing and data visuals",
        "all experiments may become incomparable, the report may be irreproducible, and presentation quality may fall if data versions or document sections conflict.",
        bullet_id,
    )

    add_work_package(
        doc,
        2,
        "Sentiment modelling",
        "Own the complete sentiment-analysis research stream from baseline to deployable model.",
        "Someone comfortable with scikit-learn, text features, Transformer fine-tuning and controlled ML experiments. Workload is concentrated in the modelling phase.",
        [
            "Implement TF-IDF word and character n-gram baselines with Logistic Regression and/or SVM.",
            "Fine-tune a suitable Transformer sentiment classifier on the frozen dataset.",
            "Run hyperparameter, preprocessing and feature ablation experiments.",
            "Save the selected model and export test predictions using the agreed schema.",
            "Document reproducible commands, configuration, seeds and model limitations.",
        ],
        "Sentiment training code, experiment records, baseline and Transformer results, a saved inference model and a standard prediction file.",
        [
            "A majority baseline, TF-IDF word/character model and one justified Transformer are reproducible.",
            "All experiments use dataset_v1, documented preprocessing, fixed configurations and multiple seeds or confidence intervals.",
            "At least one useful ablation isolates a preprocessing, feature or modelling decision.",
            "Overall and per-variety prediction files follow the agreed schema and the saved model works through inference.",
            "Results include limitations and do not rely on accuracy alone.",
        ],
        "Sentiment methods, model justification, ablation design, results and limitations.",
        "A2 and B1 through rigorous NLP comparison and reproducible sentiment evidence",
        "A1 and B3 through interpretation of cross-variety findings",
        "the project loses its traditional-versus-modern comparison and a major share of its technical evidence if baselines, ablations or reproducibility are weak.",
        bullet_id,
    )

    add_work_package(
        doc,
        3,
        "Sarcasm modelling",
        "Own the complete sarcasm-detection stream, with particular attention to class imbalance and uncertain cases.",
        "Someone interested in difficult classification, imbalanced data, thresholds, precision-recall trade-offs and qualitative language errors.",
        [
            "Implement a reproducible traditional NLP sarcasm baseline.",
            "Fine-tune a Transformer sarcasm classifier on the frozen dataset.",
            "Compare class weighting, sampling or focal-loss strategies where appropriate.",
            "Analyse the precision-recall trade-off and define a defensible decision threshold.",
            "Save the selected model and export test predictions using the agreed schema.",
        ],
        "Sarcasm training code, imbalance experiments, threshold analysis, a saved inference model and a standard prediction file.",
        [
            "A traditional baseline and one justified Transformer are reproducible on dataset_v1.",
            "Class imbalance is quantified and at least two defensible treatments or thresholds are compared.",
            "Macro-F1, per-class precision/recall and per-variety results are reported with repeated-run uncertainty where feasible.",
            "Low-confidence, context-dependent and sarcastic failure cases are exported for evaluation and application examples.",
            "The selected model artefact and prediction file work through the agreed inference contract.",
        ],
        "Sarcasm methods, imbalance strategy, quantitative results, error patterns and task limitations.",
        "A2 and B1 through technically difficult classification and imbalance handling",
        "B2 and B3 through useful uncertainty and error communication",
        "the product may make overconfident or misleading claims and lose a central user-facing feature if sarcasm evaluation is superficial.",
        bullet_id,
    )

    add_work_package(
        doc,
        4,
        "Variety-aware innovation and cross-variety module",
        "Provide the distinctive research contribution that separates VarietyLens from a standard sentiment classifier.",
        "Someone who enjoys research design, ablations, cross-domain experiments and interpreting negative as well as positive results. This module has the greatest experimental uncertainty.",
        [
            "Define a variety-agnostic control model and one variety-aware approach using known variety or market information.",
            "Implement a defensible strategy such as variety tokens, conditioned heads or a transparent routing design.",
            "Run cross-variety train/test experiments and produce a generalisation matrix.",
            "Analyse whether performance differences indicate a reliability or fairness concern.",
            "Treat automatic variety detection as optional and avoid presenting language variety as nationality or personal identity.",
        ],
        "Variety model or router, variety-aware experiment, cross-variety prediction files and clearly stated innovation findings.",
        [
            "The control and variety-aware systems differ in only the intended experimental factor wherever possible.",
            "Ablation, cross-variety and per-variety results use the common evaluation protocol.",
            "Multiple seeds or confidence intervals show whether observed differences are stable.",
            "Positive, neutral and negative findings are interpreted without claiming that the model understands culture or identifies nationality.",
            "The application receives a clear model-comparison output and a human-review condition when systems disagree.",
        ],
        "Innovation, English-variety methodology, cross-variety generalisation and implications.",
        "A1 through distinctive project design, A2 through controlled evaluation, and B1 through advanced method justification",
        "B2 and B3 through responsible comparison and original interpretation",
        "the topic may look like a standard sentiment/sarcasm classifier and lose its strongest innovation argument if the control, ablation or cross-variety evidence is missing.",
        bullet_id,
    )

    add_work_package(
        doc,
        5,
        "Evaluation, error analysis and figures",
        "Evaluate all modelling streams consistently without changing their training implementations.",
        "Someone strong in statistics, metrics, reproducible analysis, data visualisation and careful interpretation. Workload increases after model handoff.",
        [
            "Build one evaluation script that reads standard prediction files rather than importing model code.",
            "Report Macro-F1, precision, recall and confusion matrices overall and by English variety.",
            "Analyse performance by source domain and, where useful, model confidence or calibration.",
            "Create the final report and poster figures using a consistent visual language.",
            "Conduct qualitative error analysis and identify limitations, failure cases and recommendations.",
        ],
        "Reusable evaluation code, verified result tables, charts, calibration or confidence analysis, and an error-analysis dataset.",
        [
            "The evaluator consumes standard prediction files and reproduces every reported number.",
            "Macro-F1, per-class, per-variety, per-source, cross-variety and confusion-matrix results are complete.",
            "Repeated-run variability or bootstrap confidence intervals and calibration/confidence evidence are reported where feasible.",
            "A structured analysis covers approximately 40-60 representative errors without exposing private content.",
            "One locked results package supplies identical numbers and figures to the report, application and poster.",
        ],
        "Evaluation protocol, findings, error analysis, discussion, ethics, limitations and recommendations.",
        "A2 through evaluation quality, A3 through evidence visuals, B1 through reproducible metrics, and B3 through interpretation",
        "B2 through confidence and dashboard insights",
        "technical claims may become unsupported or inconsistent across deliverables if metrics, uncertainty, errors and final figures are not centrally controlled.",
        bullet_id,
    )

    add_work_package(
        doc,
        6,
        "Web application, deployment and user experience",
        "Deliver the end-user application independently from the training notebooks and expose model results through a stable inference interface.",
        "Someone comfortable with frontend or Streamlit development, integration, accessible interaction design, deployment, user testing and visual poster assembly. Workload continues throughout development and peaks near integration.",
        [
            "Build the Streamlit or equivalent web application using a mock predictor before final models are ready.",
            "Implement single-text and CSV batch input, known market/variety selection, sentiment prediction and sarcasm prediction.",
            "Display confidence, generic-versus-variety-aware comparison, human-review flags and per-variety performance in an understandable dashboard.",
            "Integrate final model artefacts through one inference wrapper without modifying their training code.",
            "Provide filters, text-level drill-down, clear upload errors, plain-language explanations, export, privacy messaging and local/deployed run instructions.",
            "Conduct the defined cross-region usability study, implement at least one evidence-based interface revision, and produce the final poster layout from team-supplied content.",
        ],
        "Runnable web application, inference wrapper, deployment instructions, usability findings, screenshots and final poster source file.",
        [
            "The app completes the full workflow from input to sentiment, sarcasm, confidence, comparison, review flag, dashboard and export.",
            "A mock predictor allows early development; final artefacts replace it without redesigning the interface.",
            "Approximately 6-12 users familiar with en-AU, en-IN or en-UK complete common tasks plus regional, sarcastic and ambiguous inputs; results remain qualitative if the sample is small.",
            "Task completion, difficulty, assistance, clarity ratings and open feedback are recorded, followed by at least one documented design improvement.",
            "The app handles invalid files and low-confidence cases, avoids identity claims, explains data handling and has both a deployment path and a local demonstration backup.",
            "The poster remains readable, uses consistent visual encoding and contains only verified claims and figures.",
        ],
        "Application design, user interaction, system integration, usability evaluation and deployment.",
        "A3 through poster communication and B2 through design features, ease of use and communication of insights",
        "A1, A2 and B3 through the visible product, system integration and application narrative",
        "the project may be judged as a model demo rather than a real end-user application if batch workflow, uncertainty, usability evidence or polished communication is missing.",
        bullet_id,
    )

    doc.add_page_break()
    add_heading(doc, "5. Interface agreements", 1)
    add_body(
        doc,
        "These agreements should be confirmed before full development begins. They are the main mechanism for parallel work: members exchange stable files and interfaces instead of editing one another's modules.",
    )

    add_heading(doc, "5.1 Frozen dataset schema", 2)
    add_schema_table(doc)
    add_body(
        doc,
        "Rule: all modelling and evaluation modules must use the same stable id and split values. Any change to dataset_v1 must produce a new version and a short reason for the change.",
        bold_prefix="Rule:",
        after=8,
    )

    add_heading(doc, "5.2 Prediction-file schema", 2)
    add_prediction_table(doc)
    add_body(
        doc,
        "Rule: Member 5 evaluates submitted prediction files. The evaluation module should not depend on internal training classes or notebooks.",
        bold_prefix="Rule:",
        after=8,
    )

    add_heading(doc, "5.3 Web inference response", 2)
    add_body(doc, "Members 2-4 should expose their final artefacts through a predictable response structure:")
    add_code_block(
        doc,
        '{\n'
        '  "variety": {"label": "en-AU", "confidence": 0.82},\n'
        '  "sentiment": {"label": "positive", "confidence": 0.91},\n'
        '  "sarcasm": {"label": "non-sarcastic", "confidence": 0.73}\n'
        '}',
    )
    add_body(
        doc,
        "Rule: Member 6 may develop the application with a mock implementation of this response. Replacing the mock with final artefacts should not require a redesign of the user interface.",
        bold_prefix="Rule:",
    )

    add_heading(doc, "6. File ownership and repository layout", 1)
    add_body(
        doc,
        "One member owns each high-change directory. Other members consume outputs from that directory but do not make routine edits inside it. This reduces merge conflicts and ambiguous responsibility.",
    )
    add_code_block(
        doc,
        "data/processed/              Member 1\n"
        "src/data/                    Member 1\n"
        "src/models/sentiment/        Member 2\n"
        "src/models/sarcasm/          Member 3\n"
        "src/models/variety/          Member 4\n"
        "evaluation/                  Member 5\n"
        "figures/                     Member 5\n"
        "app/                         Member 6\n"
        "src/inference/               Member 6\n"
        "poster/                      Member 6\n"
        "report/member_1/ ... /6/     Respective member",
    )
    add_callout(
        doc,
        "OWNERSHIP RULE",
        "Ownership means responsibility for the working state of a module. It does not prevent normal questions, demonstrations or voluntary assistance, but it avoids multiple people independently rewriting the same files.",
    )

    add_heading(doc, "7. Parallel workflow", 1)
    steps = [
        "Agree on the problem statement, target users, minimum viable product and the three interface contracts.",
        "Member 1 publishes a small schema-valid sample immediately and freezes dataset_v1 within two or three working days.",
        "Members 2-4 develop modelling streams in parallel; Member 5 builds evaluation against mock prediction files; Member 6 builds the application against a mock inference response.",
        "Model owners export versioned artefacts and prediction files. Member 5 runs the common evaluation without changing model code.",
        "Member 6 replaces mock inference with final artefacts and records usability observations. Member 5 supplies verified figures and headline findings.",
        "Each member completes their report section. Member 1 assembles the report, while Member 6 assembles the poster from approved text and figures.",
    ]
    for step in steps:
        add_list_item(doc, step, number_id)

    add_heading(doc, "8. Work that must remain shared", 1)
    add_body(doc, "The goal is low interference, not zero communication. The team must still agree on:")
    for item in [
        "The final research question, target users and project scope.",
        "The frozen dataset, prediction schema and inference response.",
        "Which experimental results are valid enough to report.",
        "The final claims, limitations and recommendations.",
        "The submitted report, poster, code package and contribution statement.",
    ]:
        add_list_item(doc, item, bullet_id)
    add_body(
        doc,
        "These decisions can be handled through short milestone meetings. They do not require a permanent cross-review assignment.",
        italic=True,
        color=MUTED,
    )

    add_heading(doc, "9. SPARKPlus and contribution evidence", 1)
    add_body(
        doc,
        "SPARKPlus should reflect actual participation, reliability, communication and delivery. It is primarily a mechanism for recognising contribution differences; it does not require a separate peer-review work package.",
    )
    add_body(doc, "Each member should retain evidence such as:")
    for item in [
        "Assigned issues or task-board records.",
        "Commits, experiment configurations and generated artefacts.",
        "Meeting attendance and short decision records.",
        "Completed report text, figures or application features.",
        "Evidence of meeting agreed deadlines and helping resolve integration issues when needed.",
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "10. Cross-region usability study", 1)
    add_body(
        doc,
        "The usability study supports B2 by testing whether people can complete the analyst workflow and understand sentiment, sarcasm, confidence, model disagreement and human-review flags. It does not replace model evaluation on the labelled BESSTIE test set.",
    )
    add_heading(doc, "10.1 Participants", 2)
    for item in [
        "Preferred: approximately 9-12 participants, with 3-4 people familiar with each of en-AU, en-IN and en-UK.",
        "Minimum practical classroom study: approximately 6 participants, with 2 people familiar with each variety.",
        "Record self-reported familiarity with an English variety, not inferred nationality or identity. Treat regional comparisons as qualitative when samples are small.",
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "10.2 Input design", 2)
    add_body(doc, "Every participant should complete the same common tasks before contributing local examples:")
    for item in [
        "Use a fixed set of 8-12 demonstration texts covering clear positive, clear negative, obvious sarcasm, mixed sentiment, polite understatement, context dependence, model disagreement and low confidence.",
        "Enter or paraphrase one short public-style expression familiar in the English variety they know.",
        "Enter or paraphrase one realistic sarcastic complaint.",
        "Enter one deliberately ambiguous statement that should encourage uncertainty or human review.",
        "Upload the same 15-30 row sample CSV, filter by variety, sentiment, sarcasm, confidence and model disagreement, then export results.",
        "Do not paste private messages, names, identifiable customer content, confidential workplace data or other personal information.",
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "10.3 Evidence to collect", 2)
    for item in [
        "Task completion, completion difficulty or time, errors and whether assistance was required.",
        "Clarity ratings for sentiment, sarcasm, confidence, model comparison, review flags and dashboard insights.",
        "Open feedback on the most confusing and most useful features.",
        "At least one documented sequence of user observation, design change and resulting improvement.",
    ]:
        add_list_item(doc, item, bullet_id)
    add_callout(
        doc,
        "PRIVACY AND VALIDITY",
        "Participant-written texts may support qualitative usability and interpretation findings. They should not be treated as a statistically representative model benchmark unless a separate consented annotation and evaluation design is completed.",
    )

    add_heading(doc, "11. Project quality gates", 1)
    gates = [
        ("Gate 1 - Scope and data freeze", "Primary user, problem, three research questions, dataset_v1, schemas and data limitations are agreed."),
        ("Gate 2 - Baseline evidence", "Sentiment and sarcasm baselines run; the common evaluator reads mock predictions; the application runs against a mock response."),
        ("Gate 3 - Advanced modelling", "Transformers, variety-aware control/ablation, repeated-run uncertainty and cross-variety experiments are complete."),
        ("Gate 4 - Product evidence", "End-to-end inference, batch workflow, filters, uncertainty, export, user testing and one interface revision are complete."),
        ("Gate 5 - Submission consistency", "Clean-environment reproduction succeeds and report, code, application and poster use the same locked results and accurate contributions."),
    ]
    for label, definition in gates:
        add_body(doc, f"{label}: {definition}", bold_prefix=f"{label}:", after=5)

    doc.add_page_break()
    add_heading(doc, "12. Final contribution statement templates", 1)
    add_body(
        doc,
        "Replace [Name] and revise each statement to match work that was actually completed. Do not claim planned work that was not delivered.",
        italic=True,
        color=MUTED,
    )

    statements = [
        (
            "Member 1",
            "[Name] developed the BESSTIE data acquisition and preprocessing pipeline, conducted exploratory data analysis, checked duplicates and potential data leakage, and created the frozen project splits and reusable data loader. They maintained the agreed repository conventions, documented the dataset and project workflow, and assembled the final report from the team-authored sections.",
        ),
        (
            "Member 2",
            "[Name] implemented and tuned the sentiment-analysis baselines and Transformer model, conducted feature, preprocessing and hyperparameter experiments, and produced the reproducible sentiment results and deployable model artefact. They authored the sentiment methodology, results and limitations sections.",
        ),
        (
            "Member 3",
            "[Name] implemented and evaluated the sarcasm-detection baselines and Transformer model, investigated class-imbalance and threshold strategies, and produced the reproducible sarcasm results and deployable model artefact. They authored the sarcasm methodology, error patterns and limitations sections.",
        ),
        (
            "Member 4",
            "[Name] developed the variety-aware modelling component and its variety-agnostic control, conducted ablation and cross-variety generalisation experiments, and analysed reliability differences across Australian, Indian and British English. They authored the project's innovation and cross-variety analysis sections.",
        ),
        (
            "Member 5",
            "[Name] designed and implemented the common evaluation pipeline, generated overall and per-variety metrics, confusion matrices, calibration or confidence analysis, and qualitative error analysis. They produced the final result figures and authored the evaluation, findings, discussion, ethics, limitations and recommendations sections.",
        ),
        (
            "Member 6",
            "[Name] designed and implemented the VarietyLens web application, including single-text and CSV input, English-variety control, sentiment and sarcasm results, model comparison, confidence and human-review displays, filtering, dashboard insights and export. They integrated the final model artefacts, conducted the cross-region usability study, implemented evidence-based interface improvements, documented deployment, authored the application sections and assembled the final poster from team-supplied content.",
        ),
    ]
    for label, text in statements:
        add_heading(doc, label, 2)
        p = add_body(doc, text, after=9)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)

    add_heading(doc, "13. Team agreement checklist", 1)
    checklist = [
        "Names have been assigned to all six work packages.",
        "The minimum viable product and optional features are separated.",
        "dataset_v1, prediction-file and web-inference formats are confirmed.",
        "Each member knows their directory, artefacts and report section.",
        "Model and result filenames include a version identifier.",
        "The team has agreed on milestone dates for data freeze, model handoff, application integration, report assembly and poster completion.",
        "Contribution statements will be updated from actual evidence before submission.",
    ]
    for item in checklist:
        add_list_item(doc, item, bullet_id)

    add_callout(
        doc,
        "FINAL REMINDER",
        "Low-dependency work succeeds only when interfaces are agreed early and kept stable. Each member should optimise their own module while protecting the handoff formats used by the rest of the team.",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
